"""Read back the saved workbook and exact append-only registry batch."""
from copy import copy
from pathlib import Path
import json
import xml.etree.ElementTree as ET

from openpyxl import load_workbook

from scripts.artifact_io import load_json_object, sha256_file, write_json_exclusive
from scripts.registry_connection import open_registry

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "outputs/adsorption_table_06_10_sync_20261022_v1"
DB = ROOT / "data/project_registry.sqlite3"
PREFIX = "fe110_c2_06_10_adsorption_20261013_"


def cell_snapshot(cell):
    return (cell.value, cell.data_type, cell.number_format,
            copy(cell.font), copy(cell.fill), copy(cell.border),
            copy(cell.alignment), copy(cell.protection))


def assert_preserved(before, after, rows, columns):
    for row in range(1, rows + 1):
        for col in range(1, columns + 1):
            assert cell_snapshot(before.cell(row, col)) == cell_snapshot(after.cell(row, col)), (before.title, row, col)
    assert str(before.merged_cells) == str(after.merged_cells)
    assert before.freeze_panes == after.freeze_panes
    for name, dimension in before.column_dimensions.items():
        assert dimension.width == after.column_dimensions[name].width, name
    for number, dimension in before.row_dimensions.items():
        assert dimension.height == after.row_dimensions[number].height, number


def main():
    manifest = load_json_object(OUT / "sheet_manifest.json")
    receipt = load_json_object(OUT / "view_receipt.json")
    book = Path(manifest["workbook"])
    assert sha256_file(book) == receipt["workbook_sha256_after"]
    assert sha256_file(OUT / "before_sync.xlsx") == manifest["workbook_sha256_before"]
    assert sha256_file(OUT / "sheet_manifest.json") == receipt["sheet_manifest_sha256"]
    original = load_workbook(OUT / "before_sync.xlsx", data_only=False)
    current = load_workbook(book, data_only=False)
    assert original.sheetnames == current.sheetnames
    assert current["最终吸附数据"].max_row == 64
    assert_preserved(original["最终吸附数据"], current["最终吸附数据"], 64, 8)
    assert_preserved(original["已登记待接受"], current["已登记待接受"], 7, 6)
    pending = current["已登记待接受"]
    actual = list(pending.iter_rows(min_row=5, max_row=19, max_col=6, values_only=True))
    assert pending.max_row == 19
    assert [row[0] for row in actual] == manifest["all15_names"]
    assert len({row[5] for row in actual}) == 15
    assert [list(row) for row in actual[3:]] == manifest["append_rows"]
    assert all(isinstance(value, str) for row in actual for value in row)
    original.close()
    current.close()

    batch = load_json_object(OUT / "registry_batch.json")
    matched = 0
    with open_registry(DB) as connection:
        for table, entries in batch["rows"].items():
            primary = [row["name"] for row in connection.execute(f'PRAGMA table_info("{table}")') if row["pk"]]
            assert len(primary) == 1, table
            for entry in entries:
                stored = connection.execute(f'SELECT * FROM "{table}" WHERE "{primary[0]}"=?', (entry[primary[0]],)).fetchone()
                assert stored is not None, (table, entry[primary[0]])
                for column, value in entry.items():
                    assert stored[column] == value, (table, entry[primary[0]], column)
                matched += 1
        for change in batch["workflow_status_changes"]:
            stored = connection.execute("SELECT * FROM calculation_workflow_status_history WHERE status_change_id=?", (change["status_change_id"],)).fetchone()
            assert stored is not None
            for key, value in change.items():
                column = "previous_workflow_status" if key == "expected_workflow_status" else key
                assert stored[column] == value, (change["status_change_id"], key)
        statuses = {}
        for name in manifest["all15_names"]:
            cid = PREFIX + name
            status = connection.execute("SELECT workflow_status FROM calculations WHERE calculation_id=?", (cid,)).fetchone()[0]
            statuses[name] = status
            assert status == ("duplicate" if name in {"06_C", "09_B"} else "needs_review"), (name, status)
            results = list(connection.execute("SELECT * FROM results WHERE calculation_id=?", (cid,)))
            assert len(results) == 3, (name, len(results))
            assert all(not row["validation_status"].startswith("accepted") for row in results), name
            if name in {"07_C", "10_C"}:
                assert "last_toten" in {row["result_name"] for row in results}
                assert "final_toten" not in {row["result_name"] for row in results}
            if name == "09_B":
                identity = next(row for row in results if row["result_name"] == "final_state_identity")
                assert json.loads(identity["text_value"])["duplicate_of"] == PREFIX + "07_A"
        assert connection.execute("PRAGMA quick_check").fetchone()[0] == "ok"
        assert not list(connection.execute("PRAGMA foreign_key_check"))
    suites = ET.parse(OUT / "tests.xml").getroot()
    test_count = sum(1 for _ in suites.iter("testcase"))
    assert not list(suites.iter("failure")) and not list(suites.iter("error"))
    assert matched == 168
    write_json_exclusive(OUT / "verification.json", {
        "status": "PASS", "formal_sheet_rows_preserved": 64,
        "existing_pending_rows_preserved": 7, "pending06_10_records": 15,
        "metadata_rows_added": 12, "registry_rows_read_back": matched,
        "workflow_transitions_read_back": len(batch["workflow_status_changes"]),
        "unaccepted_numeric_energies_published": 0, "workflow_statuses": statuses,
        "sqlite_quick_check": "ok", "foreign_key_violations": 0,
        "scoped_tests_passed": test_count, "workbook_sha256": sha256_file(book),
    })
    print(f"PASS:15 table records;168 registry rows;64 formal rows unchanged;{test_count} tests")


if __name__ == "__main__":
    main()
