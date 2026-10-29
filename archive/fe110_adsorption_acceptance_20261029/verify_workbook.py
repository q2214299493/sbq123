"""Independent saved-XLSX readback using bundled Python; no project imports."""
from copy import copy
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "outputs/adsorption_acceptance_20261029_v1"


def same_cell(first, second, *, value=True):
    if value and (first.value != second.value or first.data_type != second.data_type):
        return False
    for field in ("font", "fill", "border", "alignment", "protection"):
        if copy(getattr(first, field)) != copy(getattr(second, field)):
            return False
    return first.number_format == second.number_format


def check_preserved_sheet(first, second, row_count, columns, changed):
    errors = []
    for row in first.iter_rows(min_row=1, max_row=row_count, min_col=1, max_col=len(columns)):
        for cell in row:
            if not same_cell(cell, second[cell.coordinate], value=cell.coordinate not in changed):
                errors.append(first.title + "!" + cell.coordinate + " changed outside the authorized values")
    if set(map(str, first.merged_cells.ranges)) != set(map(str, second.merged_cells.ranges)) or first.freeze_panes != second.freeze_panes:
        errors.append(first.title + " merges/freeze panes changed")
    for col in columns:
        a, b = first.column_dimensions[col], second.column_dimensions[col]
        if (a.width, a.hidden, a.outlineLevel) != (b.width, b.hidden, b.outlineLevel):
            errors.append(first.title + " column dimension changed: " + col)
    for row in range(1, row_count + 1):
        a, b = first.row_dimensions[row], second.row_dimensions[row]
        if (a.height, a.hidden, a.outlineLevel) != (b.height, b.hidden, b.outlineLevel):
            errors.append(first.title + " row dimension changed: " + str(row))
    return errors


def main():
    manifest = json.loads((OUT / "publication_manifest.json").read_text(encoding="utf-8-sig"))
    before = load_workbook(OUT / "before.xlsx")
    after = load_workbook(manifest["workbook"])
    errors = []
    if before.sheetnames != after.sheetnames:
        errors.append("sheet names/order changed")
    for sheet_name, row_count, columns in (("最终吸附数据", 64, "ABCDEFGH"), ("已登记待接受", 19, "ABCDEF")):
        changed = {"E" + str(row["row"]) for row in manifest["pending_updates"]} if sheet_name == "已登记待接受" else set()
        errors.extend(check_preserved_sheet(before[sheet_name], after[sheet_name], row_count, columns, changed))
    main_sheet = after["最终吸附数据"]
    if main_sheet.max_row != 67:
        errors.append("main-sheet row count is not 67")
    for item in manifest["accepted"]:
        request = json.loads(Path(item["request"]).read_text(encoding="utf-8-sig"))
        expected = [item["toten_eV"] if column["source"]["type"] == "result" else column["source"]["value"]
                    for column in request["columns"]]
        actual = [main_sheet.cell(item["row"], col).value for col in range(1, 9)]
        if actual != expected or main_sheet.cell(item["row"], 6).data_type != "n":
            errors.append("accepted source values/types differ at row " + str(item["row"]))
    pending = after["已登记待接受"]
    if pending.max_row != 23:
        errors.append("pending-sheet row count is not 23")
    for item in manifest["pending_updates"]:
        if pending.cell(item["row"], 5).value != item["reason"]:
            errors.append("pending review reason differs")
    for index, expected in enumerate(manifest["pending_append"]):
        if [pending.cell(20 + index, col).value for col in range(1, 7)] != expected:
            errors.append("pending source differs at appended row " + str(20 + index))
    if any(cell.data_type == "e" for sheet in after for row in sheet for cell in row):
        errors.append("saved workbook contains an Excel error cell")
    report = {"status": "PASS" if not errors else "FAIL", "errors": errors[:30],
        "original_main_rows_preserved": 64, "original_pending_identities_preserved": 15,
        "original_cell_styles_dimensions_merges_freeze_preserved": not errors,
        "accepted_numeric_rows_verified": 3, "pending_metadata_rows_added": 4,
        "main_row_count": main_sheet.max_row, "pending_row_count": pending.max_row,
        "workbook_sha256": hashlib.sha256(Path(manifest["workbook"]).read_bytes()).hexdigest(),
        "new_scientific_jobs": 0}
    destination = OUT / "workbook_verification.json"
    if destination.exists():
        if json.loads(destination.read_text(encoding="utf-8")) != report:
            raise ValueError("preserve earlier verification; source or result changed")
    else:
        with destination.open("x", encoding="utf-8") as stream:
            json.dump(report, stream, ensure_ascii=False, indent=2)
    print(json.dumps(report, ensure_ascii=True))
    if errors:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
