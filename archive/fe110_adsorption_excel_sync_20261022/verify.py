"""Read-only saved-XLSX/registry verification using the bundled spreadsheet runtime."""
from copy import copy
from pathlib import Path

import openpyxl

from scripts.artifact_io import load_json_object, sha256_file, sha256_json, write_json_exclusive
from scripts.registry_connection import open_registry


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "outputs/adsorption_table_sync_20261022_v1"
BOOK = ROOT / "outputs/adsorption_topic1_20260702/课题一吸附_最终.xlsx"
DB = ROOT / "data/project_registry.sqlite3"


def verify_original(before, after):
    for row in before.iter_rows():
        for cell in row:
            other = after[cell.coordinate]
            if cell.value != other.value:
                raise ValueError("original cell value/formula changed:" + cell.coordinate)
            for field in ("font", "fill", "border", "alignment", "number_format", "protection"):
                if copy(getattr(cell, field)) != copy(getattr(other, field)):
                    raise ValueError("original cell style changed:" + cell.coordinate + ":" + field)
    if list(before.merged_cells.ranges) != list(after.merged_cells.ranges) or before.freeze_panes != after.freeze_panes:
        raise ValueError("original merge/freeze changed")
    for key, dimension in before.column_dimensions.items():
        if dimension.width != after.column_dimensions[key].width:
            raise ValueError("original column width changed")
    for index, dimension in before.row_dimensions.items():
        if index <= 55 and dimension.height != after.row_dimensions[index].height:
            raise ValueError("original row height changed")
    if after.max_row != 64 or after.max_column != 8:
        raise ValueError("unexpected final workbook structure")


def verify_published(manifest, after):
    rows = []
    with open_registry(DB) as connection:
        for index, item in enumerate(manifest["accepted"]):
            request = load_json_object(Path(item["request"]))
            receipt = load_json_object(Path(request["receipt_path"]))
            plan = load_json_object(OUT / (request["promotion_id"] + ".plan.json"))
            row = 56 + index
            values = [after.cell(row, column).value for column in range(1, 9)]
            if values != plan["row_values"] or sha256_json(values) != receipt["written_values_sha256"]:
                raise ValueError("published row differs from exact accepted-source plan")
            accepted = connection.execute("SELECT * FROM results WHERE result_id=?", (item["accepted_result_id"],)).fetchone()
            if accepted["validation_status"] != "accepted_compatible_final_energy" or accepted["numeric_value"] != values[5]:
                raise ValueError("formal energy not backed by an accepted compatible result")
            count = connection.execute("SELECT count(*) FROM excel_promotions WHERE registry_id=? AND promotion_id=?",
                                       (item["calculation_id"], request["promotion_id"])).fetchone()[0]
            if count != 1:
                raise ValueError("exactly one publication receipt required")
            if row != receipt["excel_row_number"] or after.cell(row, 7).data_type != "s" or after.cell(row, 8).data_type != "s":
                raise ValueError("missing E0/Eads must not be fabricated numeric values")
            rows.append({"calculation_id": item["calculation_id"], "row": row, "toten_eV": values[5],
                         "promotion_id": request["promotion_id"]})
        check = connection.execute("PRAGMA quick_check").fetchone()[0]
        foreign = list(connection.execute("PRAGMA foreign_key_check"))
        if check != "ok" or foreign:
            raise ValueError("database integrity check failed")
    return rows, check


def main():
    manifest = load_json_object(OUT / "sync_manifest.json")
    backup = OUT / "课题一吸附_最终.before_sync.xlsx"
    if sha256_file(backup) != manifest["workbook_sha256_before"]:
        raise ValueError("backup no longer matches the reviewed source workbook")
    old, current = [openpyxl.load_workbook(p, data_only=False) for p in (backup, BOOK)]
    verify_original(old["最终吸附数据"], current["最终吸附数据"])
    if current.sheetnames != ["最终吸附数据", "已登记待接受"]:
        raise ValueError("unexpected sheet names")
    rows, check = verify_published(manifest, current["最终吸附数据"])
    pending = [[current["已登记待接受"].cell(row, col).value for col in range(1, 7)] for row in range(5, 8)]
    if pending != manifest["pending_metadata_rows"] or any(not isinstance(value, str) for row in pending for value in row):
        raise ValueError("pending view must retain exact registered text metadata only")
    repeat = load_json_object(OUT / "repeat_publication_check.json")
    if not repeat["blocked"] or repeat["request_sha256"] != sha256_file(Path(manifest["accepted"][0]["request"])):
        raise ValueError("repeat publication check missing or stale")
    current_hash = sha256_file(BOOK)
    if current_hash != load_json_object(OUT / "view_export_receipt.json")["workbook_sha256_after"]:
        raise ValueError("saved view bytes changed")
    old.close()
    current.close()
    write_json_exclusive(OUT / "verification.json", {"original_rows_preserved": 55,
        "original_values_formulas_styles_preserved": True, "accepted_rows": rows,
        "pending_metadata_rows": 3, "no_unaccepted_numeric_energy_published": True,
        "repeat_publication_blocked": True, "quick_check": check, "foreign_key_check": [],
        "workbook_sha256_after": current_hash, "new_scientific_jobs": 0})
    print("VERIFIED:55 original rows unchanged;9 accepted-source rows;3 pending metadata rows;duplicate publication blocked;DB integrity ok")


if __name__ == "__main__":
    main()
