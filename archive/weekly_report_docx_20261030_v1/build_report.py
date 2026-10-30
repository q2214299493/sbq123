"""Build this week's report by patching the retained DOCX package."""

from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree


TASK = Path(__file__).resolve().parent
ROOT = TASK.parent.parent
REFERENCE = TASK / "reference.docx"
OUTPUT = ROOT / "reports/Fe110_新模型训练与AQCat25微调周报_20261026-20261030.docx"
REFERENCE_HASH = "078e43cbece55c303f170b653c3e89dc28cd431ced627fa042825d96c69b4bf5"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PR = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"w": W}


def digest(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


assert digest(REFERENCE) == REFERENCE_HASH
assert (TASK / "artifact.md").is_file()
assert not OUTPUT.exists(), "Choose a fresh output instead of overwriting a report"
facts = json.loads((TASK / "facts.json").read_text(encoding="utf-8"))
assert facts["development_mae"] == [0.31907776910299446, 0.05564634550888187]
assert [job["steps"] for job in facts["dft_rows"]] == [75, 93, 69]

with ZipFile(REFERENCE) as source_zip:
    infos = source_zip.infolist()
    original = {info.filename: source_zip.read(info.filename) for info in infos}

source_doc = Document(REFERENCE)
tree = etree.fromstring(original["word/document.xml"])
body = tree.find(qn("w:body"))
section = deepcopy(body.find(qn("w:sectPr")))
section_bytes = etree.tostring(section, method="c14n", exclusive=True)
prototypes = {
    role: deepcopy(source_doc.paragraphs[index]._p)
    for role, index in {"title": 0, "subtitle": 1, "heading": 3, "body": 4}.items()
}
table_prototype = deepcopy(source_doc.tables[0]._tbl)
for item in list(body):
    body.remove(item)
body.append(section)


def prop(parent, name, **attributes):
    element = parent.find(qn(name))
    if element is None:
        element = OxmlElement(name)
        parent.append(element)
    for key, value in attributes.items():
        element.set(qn(key), str(value))
    return element


def paragraph(text="", role="body", *, bold=False, new_page=False):
    p = deepcopy(prototypes[role])
    p.attrib.clear()
    ppr = p.find(qn("w:pPr"))
    for item in list(p):
        if item is not ppr:
            p.remove(item)
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        p.insert(0, ppr)
    prop(ppr, "w:widowControl", **{"w:val": 1})
    prop(ppr, "w:keepNext", **{"w:val": int(role in ("title", "heading"))})
    if new_page:
        prop(ppr, "w:pageBreakBefore", **{"w:val": 1})
    source_run = prototypes[role].find(qn("w:r"))
    source_rpr = source_run.find(qn("w:rPr")) if source_run is not None else None
    run = OxmlElement("w:r")
    rpr = deepcopy(source_rpr) if source_rpr is not None else OxmlElement("w:rPr")
    prop(rpr, "w:color", **{"w:val": "000000"})
    if bold:
        prop(rpr, "w:b", **{"w:val": 1})
    run.append(rpr)
    text_node = OxmlElement("w:t")
    text_node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    text_node.text = text
    run.append(text_node)
    p.append(run)
    return p


def add(text="", role="body", **kwargs):
    p = paragraph(text, role, **kwargs)
    body.insert(len(body) - 1, p)
    return p


def table(rows, widths):
    result = deepcopy(table_prototype)
    for row in list(result.findall(qn("w:tr"))):
        result.remove(row)
    grid = result.find(qn("w:tblGrid"))
    for column in list(grid):
        grid.remove(column)
    for width in widths:
        prop_column = OxmlElement("w:gridCol")
        prop_column.set(qn("w:w"), str(width))
        grid.append(prop_column)
    tpr = result.find(qn("w:tblPr"))
    prop(tpr, "w:tblW", **{"w:type": "dxa", "w:w": sum(widths)})
    prop(tpr, "w:tblLayout", **{"w:type": "fixed"})
    for i, values in enumerate(rows):
        tr = OxmlElement("w:tr")
        trpr = OxmlElement("w:trPr")
        trpr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        tr.append(trpr)
        for j, text in enumerate(values):
            tc = OxmlElement("w:tc")
            tcpr = OxmlElement("w:tcPr")
            prop(tcpr, "w:tcW", **{"w:type": "dxa", "w:w": widths[j]})
            prop(tcpr, "w:vAlign", **{"w:val": "center"})
            if i == 0:
                prop(tcpr, "w:shd", **{"w:val": "clear", "w:fill": "E8EEF5"})
            tc.append(tcpr)
            p = paragraph(str(text), bold=(i == 0))
            ppr = p.find(qn("w:pPr"))
            prop(ppr, "w:spacing", **{"w:before": 0, "w:after": 0, "w:line": 252, "w:lineRule": "auto"})
            prop(ppr, "w:jc", **{"w:val": "left" if j == 0 and i else "center"})
            tc.append(p)
            tr.append(tc)
        result.append(tr)
    body.insert(len(body) - 1, result)
    spacer = add("")
    prop(spacer.find(qn("w:pPr")), "w:spacing", **{"w:after": 0, "w:line": 80, "w:lineRule": "exact"})


add("Fe110 新模型训练与微调周报", "title")
add("NequIP 从头训练与 AQCat25 微调", "subtitle")
add("汇报周期：2026年10月26日至10月30日")
add("结果更新至10月29日已保存的训练与计算检查点")

add("完成内容", "heading")
add("本周完成了512帧力训练数据包、NequIP新模型从头训练及AQCat25（AQC25）限时微调对照。微调候选在同一开发集上的力误差较低，一例C₂H₃经微调预松弛后接DFT，达到0.02 eV/Å力标准。整体DFT加速比与跨体系泛化仍待验证。[1–3]")

add("训练数据扩充", "heading")
add("训练数据由原366帧扩充至512帧，新增119个已审核原选点单点，以及27个定向补充帧（11个单点、两条优化轨迹各8帧）。开发集78帧、保留测试53帧，另16个历史观察留出继续隔离，分区不重新划分。[1]")
add("训练只拟合力，能量损失系数为0。底部18个固定Fe不参与损失，其原始力保留；其余可动原子均匀加权，训练帧均匀有放回采样。逐项核对659个清单样本及590行训练与开发数据库，坐标、原始力和分区等检查通过。磁性等级仍分别为258帧局域正Fe审核、248帧旧来源审核、6帧仅总磁矩审核，未提升其审核等级。[1]")

add("数据覆盖检查", "heading")
add("以固定1597帧训练候选池检查512帧数据，组成、局域环境与受力区间的联合类型缺口由189类降至178类，新增覆盖11类；SOAP固定阈值下仍有4个超阈值候选帧。[4]")
table([
    ["组成", "训练帧数", "≤0.05 eV/Å帧数", "剩余联合类型"],
    ["C₂HO", 122, 9, 20],
    ["C₂H₂O", 190, 17, 25],
    ["C₂H₃", 51, 4, 39],
    ["C₃H₂O", 126, 0, 94],
], [2448, 2016, 3024, 2304])
add("表中489帧为四种目标组成，另有23帧回放数据。低力指标取可动原子最大力。C₂H₃的4个低力帧仍只来自10_B同一来源；C₃H₂O最小值约0.09057 eV/Å，仍缺≤0.05 eV/Å标签。178类是覆盖诊断结果，不等于需要新增178份DFT计算。[4]")

add("新模型训练与微调对比", "heading")
add("两条路线使用同一512帧训练集和78帧开发集，在MZ73 GPU1顺序运行，单seed为42，每项预算30分钟。开发选模均采用原角色加权可动力向量MAE。[2]")
table([
    ["项目", "NequIP 从头训练", "AQCat25 微调"],
    ["初始化", "随机初始化", "原始AQCat25权重"],
    ["模型设置", "3层 l_max=1\n32 features 约21.6万参数", "保留原模型架构"],
    ["学习率与批量", "0.001；batch 4", "0.00001；batch 1"],
    ["计划轮数上限", "100轮", "20轮"],
    ["实际退出", "生产者正常退出", "30分钟预算触发中止"],
    ["可确认保存位置", "epoch 33\nglobal_step 4352", "最后完整断点epoch 9\nstep 4608"],
    ["开发力MAE eV/Å", "0.31908", "0.05565"],
    ["回执跨度 秒", "902.19", "1802.07"],
], [2448, 3672, 3672])
add("NequIP最佳与最后快照均保存于epoch 33，但最终训练轮数及具体停止原因尚未确认。微调未完成计划20轮；最佳候选未记录epoch，不能将最后完整断点的epoch 9当作最佳候选轮数。两项总墙钟约45分11秒，均有实际CUDA设备记录，回收23份文件的哈希检查通过。[2]")
add("微调候选的开发指标优于本次从头训练候选。该指标来自参与选模的开发集；两路线的架构、参数规模、学习率、批量和实际训练量不同，因此本轮结果不能单独证明预训练的因果贡献，也不证明从头训练的充分收敛。[2]")

add("预松弛验证", "heading")
table([
    ["路线", "体系", "ML步数", "最终ML最大力 eV/Å"],
    ["NequIP", "C₂H₃", 34, "0.09423"],
    ["AQCat25 微调", "C₂H₃", 60, "0.06936"],
    ["NequIP", "C₃H₂O", 41, "0.09442"],
    ["AQCat25 微调", "C₃H₂O", 65, "0.07734"],
], [2880, 1728, 1440, 3744])
add("4项均达到ML的0.10 eV/Å标准。C₂H₃保持原分子连接；两条C₃H₂O路线均形成CO与C₂H₂片段，不能记作完整C₃吸附态优化成功，也不能由ML力阈值推定DFT已收敛。[5]")
add("进一步用冻结微调候选完成11–15批次15个起点的ML预松弛，全部保留目标分子连接、达到ML力标准并通过端点审核。后续DFT已提交且均采用NSW 300；截至10月29日15:23北京时间的记录为2项RUN、13项PEND，最终收敛与几何尚待复核。[6]")

add("DFT验证结果", "heading")
add("C₂H₃的10_A同起点比较采用纯DFT、NequIP预松弛后DFT、微调预松弛后DFT三条路线，保持相同DFT方法与0.02 eV/Å力标准。截至10月29日22:40的终态记录如下；SCF数只列完整离子步内的电子迭代。[3]")
table([
    ["路线", "状态", "完整离子步", "完整步SCF", "DFT最大力 eV/Å"],
    ["纯DFT", "用户取消\n未收敛", 75, 1497, "0.401579"],
    ["NequIP后DFT", "用户取消\n未收敛", 93, 1212, "0.061108"],
    ["微调后DFT", "正常结束\n达到力标准", 69, 586, "0.018306"],
], [2448, 2160, 1584, 1584, 2016])
add("微调路线用69个离子步、586次SCF达到力标准，VASP成本约346.11核时。另两项按用户要求在未收敛时取消，记录为截断对照，因此无法计算三条路线到相同最低点的完整成本比；核时也尚未计入训练及DFT标注摊销。[3]")
add("三份终态或停止结构均保留C₂H₃连接和主要Fe近邻，但几何并不完全一致。两项对照未收敛，不能将其位置差异解释为三个不同的稳定最低点。全轨迹电子与磁性审核、完整科学接受登记尚未完成。[3]")

add("未验证内容与风险", "heading")
add("本轮未完成独立测试或多随机种子复验。开发集用于选模；保留测试中的部分来源已有历史评估记录，不能将其称为重新获得的全盲测试。开发误差降低也尚未证明低力区的力方向及可动Fe预测足够准确。[1–2]")
add("C₃H₂O近平衡低力标签仍缺，且两条ML试验均发生碎裂；微调仅完成限时部分训练。当前候选的适用范围和进一步训练收益尚未确定。两模型均为force-only，不能用其能量输出报告跨组成吸附能。[2,4–5]")
add("15项扩展DFT的收敛与最终化学状态未确认；现有三臂结果也未证明相同DFT最低点或整体加速，候选模型尚未正式晋级。[3,6]")

add("下周计划", "heading")
add("优先在既有训练与开发数据上，按低力区、可动Fe和吸附原子分组诊断力大小及方向误差，形成下一轮训练调整方案。随后结合15项现有DFT任务的按需检查，完成电子、力、磁性与几何审核；只有获得相同终态、相同收敛标准的完整对照后，再统计SCF、核时及标注和训练摊销成本。新增训练或计算另行审核。[2,6]")

add("主要依据", "heading")
relationships = etree.fromstring(original["word/_rels/document.xml.rels"])
evidence = [
    ("[1] 512帧训练数据包", "docs/reviews/fe110_force_package_20261027/README.md"),
    ("[2] 训练与微调终止审核", "docs/reviews/fe110_scratch512_20261028/README.md"),
    ("[3] C₂三臂DFT与结构比较", "docs/reviews/fe110_c2_three_arm_stop_20261029/README.md"),
    ("[4] 联合覆盖与低力区检查", "docs/reviews/fe110_joint_update512_20261027/README.md"),
    ("[5] 同起点ML预松弛验证", "docs/reviews/fe110_scratch512_speed2_20261028/README.md"),
    ("[6] 15项扩展DFT最新记录", "docs/reviews/fe110_c2_11_15_ft512_20261029/nsw300_update.md"),
]
for index in range(0, len(evidence), 2):
    p = add("")
    p.remove(p.find(qn("w:r")))
    for offset, (label, path) in enumerate(evidence[index:index + 2]):
        rid = f"rIdWeekly20261030Evidence{index + offset + 1}"
        target = ROOT / path
        relationship = etree.SubElement(relationships, f"{{{PR}}}Relationship")
        relationship.set("Id", rid)
        relationship.set("Type", R + "/hyperlink")
        relationship.set("Target", target.as_uri())
        relationship.set("TargetMode", "External")
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), rid)
        run = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        prop(rpr, "w:color", **{"w:val": "245A81"})
        prop(rpr, "w:u", **{"w:val": "single"})
        run.append(rpr)
        node = OxmlElement("w:t")
        node.text = label
        run.append(node)
        hyperlink.append(run)
        p.append(hyperlink)
        if offset == 0:
            spacer = OxmlElement("w:r")
            node = OxmlElement("w:t")
            node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            node.text = "    "
            spacer.append(node)
            p.append(spacer)

changed = {
    "word/document.xml": etree.tostring(tree, encoding="UTF-8", xml_declaration=True, standalone=True),
    "word/_rels/document.xml.rels": etree.tostring(relationships, encoding="UTF-8", xml_declaration=True, standalone=True),
}
for name, data in original.items():
    if name.endswith(".xml") and (name.startswith("word/header") or name.startswith("word/footer")):
        story = etree.fromstring(data)
        texts = story.findall(".//w:t", NS)
        if texts:
            texts[0].text = "Fe110 新模型训练与微调周报" if name.startswith("word/header") else "2026年10月26日至10月30日"
            for node in texts[1:]:
                node.text = ""
            changed[name] = etree.tostring(story, encoding="UTF-8", xml_declaration=True, standalone=True)
core = etree.fromstring(original["docProps/core.xml"])
for node in list(core):
    local = etree.QName(node).localname
    if local == "title":
        node.text = "Fe110 新模型训练与微调周报"
    elif local == "subject":
        node.text = "2026年10月26日至10月30日 NequIP训练与AQCat25微调"
    elif local in ("creator", "lastModifiedBy", "created", "modified"):
        core.remove(node)
changed["docProps/core.xml"] = etree.tostring(core, encoding="UTF-8", xml_declaration=True, standalone=True)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
with OUTPUT.open("xb") as stream:
    with ZipFile(stream, "w") as output_zip:
        for info in infos:
            output_zip.writestr(info, changed.get(info.filename, original[info.filename]))

with ZipFile(OUTPUT) as output_zip:
    assert all(output_zip.read(name) == data for name, data in original.items() if name not in changed)
    final_tree = etree.fromstring(output_zip.read("word/document.xml"))
    assert etree.tostring(final_tree.find("w:body/w:sectPr", NS), method="c14n", exclusive=True) == section_bytes
final = Document(OUTPUT)
assert final.paragraphs[0].style.name == "Title"
assert [len(t.rows) for t in final.tables] == [5, 9, 5, 4]
text = "\n".join(p.text for p in final.paragraphs) + "\n" + "\n".join(c.text for t in final.tables for r in t.rows for c in r.cells)
for token in ("512", "78帧", "53帧", "0.31908", "0.05565", "0.018306", "1497", "1212", "586", "用户取消", "未完成计划20轮"):
    assert token in text, token
assert "10月19日" not in text
assert digest(REFERENCE) == REFERENCE_HASH
receipt = {
    "output": str(OUTPUT), "output_sha256": digest(OUTPUT),
    "reference_sha256": REFERENCE_HASH, "section_xml_unchanged": True,
    "changed_parts": sorted(changed), "preserved_package_parts": len(original) - len(changed),
    "tables": 4, "source_links": len(evidence), "visual_QA": "PENDING",
    "source_evidence": facts["source_evidence"],
}
(TASK / "build_receipt.json").write_text(json.dumps(receipt, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(receipt, ensure_ascii=False, indent=2))
